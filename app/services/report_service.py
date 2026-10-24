"""Portable report generation. No native PDF libraries or external network fetches."""

import csv
import io
import json
from html import escape
from app.services.scoring import grade

DISCLAIMER = "Automated results cover one rendered page at one viewport. The score is a heuristic, not WCAG certification. Review evidence and test with assistive technology."


def report_data(scan):
    done = scan.status == "done"
    return {
        "schema_version": "2.0",
        "id": scan.id,
        "url": scan.url,
        "title": scan.title,
        "status": scan.status,
        "wcag_level": scan.wcag_level,
        "score": scan.score if done else None,
        "grade": grade(scan.score) if done else None,
        "principle_scores": scan.principle_scores or {},
        "issue_count": scan.issue_count,
        "severity_counts": {
            k: getattr(scan, k + "_count") for k in ("critical", "serious", "moderate", "minor")
        },
        "created_at": scan.created_at.isoformat() if scan.created_at else None,
        "completed_at": scan.completed_at.isoformat() if scan.completed_at else None,
        "notes": scan.error,
        "metadata": getattr(scan, "scan_metadata", {})
        or {"coverage": "legacy", "scoring_version": "1"},
        "issues": scan.issues or [],
        "review_items": getattr(scan, "review_items", []) or [],
        "limitations": DISCLAIMER,
    }


def json_export(scan):
    return json.dumps(report_data(scan), indent=2, ensure_ascii=False).encode("utf-8")


def _csv_safe(value):
    text = str(value if value is not None else "")
    return "'" + text if text.lstrip().startswith(("=", "+", "-", "@", "\t", "\r")) else text


def csv_export(scan):
    out = io.StringIO(newline="")
    fields = [
        "scan_id",
        "url",
        "score",
        "wcag_target",
        "coverage",
        "engine",
        "final_url",
        "scanned_at_utc",
        "notes",
        "limitations",
        "category",
        "rule_id",
        "impact",
        "message",
        "selector",
        "html",
        "wcag_criteria",
        "principle",
        "wcag_level",
        "fix",
        "help_url",
        "source",
    ]
    writer = csv.DictWriter(out, fieldnames=fields, extrasaction="ignore")
    writer.writeheader()
    d = report_data(scan)
    rows = d["issues"] + d["review_items"]
    if not rows:
        rows = [{"category": "summary", "message": "No findings detected by the checks that ran."}]
    for issue in rows:
        row = {
            **issue,
            "scan_id": scan.id,
            "url": scan.url,
            "score": d["score"],
            "wcag_target": scan.wcag_level,
            "coverage": d["metadata"].get("coverage", "legacy"),
            "engine": d["metadata"].get("engine", "Legacy engine"),
            "final_url": d["metadata"].get("final_url", scan.url),
            "scanned_at_utc": d["created_at"],
            "notes": d["notes"] or "",
            "limitations": DISCLAIMER,
            "wcag_criteria": ", ".join(issue.get("wcag_criteria") or []),
        }
        writer.writerow({k: _csv_safe(row.get(k, "")) for k in fields})
    return out.getvalue().encode("utf-8-sig")


def _sections(scan):
    d = report_data(scan)
    rows = [
        ("Website", d["url"]),
        ("Page title", d["title"] or "Untitled"),
        ("Scanned at (UTC)", d["created_at"] or "—"),
        ("WCAG target", f"2.2 Level {d['wcag_level']}"),
        ("Heuristic score", f"{d['score']}/100 — {d['grade']}"),
        ("Confirmed findings", str(d["issue_count"])),
        ("Review items", str(len(d["review_items"]))),
        ("Engine", d["metadata"].get("engine", "Legacy scan — rerun recommended")),
        ("Coverage", d["metadata"].get("coverage", "legacy")),
        ("Scan notes", d["notes"] or "None"),
    ]
    return d, rows


def pdf_export(scan):
    from reportlab.lib import colors
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.enums import TA_LEFT
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, KeepTogether
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from pathlib import Path

    output = io.BytesIO()
    styles = getSampleStyleSheet()
    # Use a portable installed Unicode font when available; standard PDF font is the fallback.
    for path in (
        "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
        "C:/Windows/Fonts/arial.ttf",
        "/System/Library/Fonts/Supplemental/Arial.ttf",
    ):
        if Path(path).is_file():
            if "AccessCheckUnicode" not in pdfmetrics.getRegisteredFontNames():
                pdfmetrics.registerFont(TTFont("AccessCheckUnicode", path))
            for name in ("Normal", "BodyText", "Title", "Heading1", "Heading2", "Heading3"):
                styles[name].fontName = "AccessCheckUnicode"
            break
    styles.add(
        ParagraphStyle(
            name="Evidence",
            parent=styles["BodyText"],
            fontSize=8,
            leading=12,
            wordWrap="CJK",
            spaceAfter=8,
        )
    )

    def para(text, style="BodyText"):
        return Paragraph(escape(str(text)).replace("\n", "<br/>"), styles[style])

    d, rows = _sections(scan)
    story = [
        para("AccessCheck", "Title"),
        para("Website accessibility report", "Heading2"),
        Spacer(1, 12),
    ]
    for key, value in rows:
        story += [para(key, "Heading3"), para(value, "Evidence")]
    story += [Spacer(1, 10), para(DISCLAIMER, "Evidence")]
    for title, items in (
        ("Confirmed findings", d["issues"]),
        ("Manual review & best practices", d["review_items"]),
    ):
        story += [para(title, "Heading1")]
        if not items:
            story.append(para("No items recorded."))
        for n, item in enumerate(items, 1):
            story += [para(f"{n}. {item.get('message', 'Finding')}", "Heading3")]
            for key in (
                "category",
                "impact",
                "rule_id",
                "wcag_criteria",
                "selector",
                "html",
                "fix",
                "help_url",
            ):
                value = item.get(key, "")
                value = ", ".join(value) if isinstance(value, list) else value
                if value:
                    story.append(para(f"{key.replace('_', ' ').title()}: {value}", "Evidence"))

    def footer(canvas, doc):
        canvas.saveState()
        canvas.setFont("Helvetica", 8)
        canvas.setFillColor(colors.HexColor("#586078"))
        canvas.drawString(42, 25, "AccessCheck • Automated accessibility report")
        canvas.drawRightString(553, 25, str(doc.page))
        canvas.restoreState()

    SimpleDocTemplate(
        output,
        pagesize=(595, 842),
        rightMargin=42,
        leftMargin=42,
        topMargin=42,
        bottomMargin=44,
        title="AccessCheck Report",
    ).build(story, onFirstPage=footer, onLaterPages=footer)
    return output.getvalue()


def docx_export(scan):
    from docx import Document
    from docx.shared import Inches, Pt, RGBColor

    doc = Document()
    sec = doc.sections[0]
    sec.header.paragraphs[0].text = "ACCESSCHECK  /  ACCESSIBILITY REPORT"
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(10)
    doc.add_heading("Website accessibility report", 0)
    d, rows = _sections(scan)
    table = doc.add_table(rows=0, cols=2)
    table.style = "Light Shading Accent 1"
    for key, value in rows:
        cells = table.add_row().cells
        cells[0].text = key
        cells[1].text = value
    doc.add_paragraph(DISCLAIMER)
    for title, items in (
        ("Confirmed findings", d["issues"]),
        ("Manual review & best practices", d["review_items"]),
    ):
        doc.add_heading(title, 1)
        if not items:
            doc.add_paragraph("No items recorded.")
        for n, item in enumerate(items, 1):
            doc.add_heading(f"{n}. {item.get('message', 'Finding')}", 2)
            for key in (
                "category",
                "impact",
                "rule_id",
                "wcag_criteria",
                "selector",
                "html",
                "fix",
                "help_url",
            ):
                value = item.get(key, "")
                value = ", ".join(value) if isinstance(value, list) else value
                if value:
                    doc.add_paragraph(f"{key.replace('_', ' ').title()}: {value}")
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


def doc_export(scan):
    """Word-compatible RTF for the legacy .doc download option."""

    def rtf(s):
        result = ""
        for ch in str(s):
            if ch in "\\{}":
                result += "\\" + ch
            elif ch == "\n":
                result += "\\line "
            elif ord(ch) > 127:
                for pos in range(0, len(b := ch.encode("utf-16-le")), 2):
                    unit = int.from_bytes(b[pos : pos + 2], "little")
                    result += f"\\u{unit if unit < 32768 else unit - 65536}?"
            else:
                result += ch
        return result

    d, rows = _sections(scan)
    lines = (
        ["AccessCheck — Website accessibility report"]
        + [f"{k}: {v}" for k, v in rows]
        + [DISCLAIMER]
    )
    for heading, items in (
        ("Confirmed findings", d["issues"]),
        ("Manual review & best practices", d["review_items"]),
    ):
        lines.append(heading)
        for item in items:
            for k, v in item.items():
                lines.append(f"{k}: {', '.join(v) if isinstance(v, list) else v}")
    return (
        "{\\rtf1\\ansi\\deff0{\\fonttbl{\\f0 Calibri;}}\\f0\\fs22 "
        + "\\par\n".join(map(rtf, lines))
        + "}"
    ).encode("ascii")
